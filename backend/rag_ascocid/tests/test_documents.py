"""Documents ajoutés par les administrateurs : lecture, intégration, retrait.

Ce qui est vérifié ici, ce sont les promesses faites à l'administrateur :
un document envoyé est trouvable par l'assistant, un document retiré ne l'est
plus nulle part (ni lexical, ni vectoriel, ni fichier), un scan est refusé
plutôt qu'indexé vide, et une reconstruction du corpus AsCoCid ne fait pas
disparaître ce qui a été ajouté à la main.
"""

from __future__ import annotations

import io

import pytest

from ascocid.application.ingestion.documents import (
    PREMIER_IDOC,
    Bibliotheque,
    DocumentRefuse,
    reinjecter,
)
from ascocid.application.requete.pipeline import Pipeline
from ascocid.domain.modeles import Corpus, Fiche, TypeFiche
from ascocid.infrastructure.graphe.sqlite import MagasinSqlite
from ascocid.infrastructure.index.memoire import IndexMemoire
from ascocid.infrastructure.parsing.document import DocumentIllisible, format_de, lire

PARAGRAPHE = (
    "La fermentation du cidre se conduit idéalement entre 10 et 14 degrés pour "
    "préserver les arômes fruités. Une température trop élevée accélère la "
    "fermentation et fait perdre une partie des composés volatils. "
)


# ── Fabrication de documents de test ────────────────────────────────────────


def pdf(pages: list[list[tuple[str, float]]]) -> bytes:
    """Un PDF dont chaque page est une liste de (texte, taille de police)."""
    import pymupdf

    document = pymupdf.open()
    for lignes in pages:
        page = document.new_page()
        y = 60
        for texte, taille in lignes:
            # insert_textbox replie les lignes longues comme le ferait un
            # traitement de texte ; le PDF garde ainsi une mise en page réaliste.
            rect = pymupdf.Rect(50, y, 545, y + 400)
            reste = page.insert_textbox(rect, texte, fontsize=taille)
            y += 400 - reste + taille if reste > 0 else 400
            y += 10
    octets = document.tobytes()
    document.close()
    return octets


def docx(blocs: list[tuple[str, str]]) -> bytes:
    """Un .docx : (« titre » | « texte » | « tableau », contenu)."""
    import docx as python_docx

    document = python_docx.Document()
    for nature, contenu in blocs:
        if nature == "titre":
            document.add_heading(contenu, level=1)
        elif nature == "tableau":
            lignes = [l.split(";") for l in contenu.split("\n")]
            table = document.add_table(rows=len(lignes), cols=len(lignes[0]))
            for i, ligne in enumerate(lignes):
                for j, cellule in enumerate(ligne):
                    table.cell(i, j).text = cellule
        else:
            document.add_paragraph(contenu)
    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


class EmbedderFactice:
    """Vecteurs déterministes, sans modèle : on teste la tuyauterie, pas E5."""

    dimension = 4
    identifiant_modele = "factice"

    def encoder(self, textes: list[str], *, requete: bool = False) -> list[list[float]]:
        return [[float(len(t) % 7 + 1), float(t.count("e") + 1), 1.0, 0.5] for t in textes]


@pytest.fixture
def espace(tmp_path):
    magasin = MagasinSqlite(tmp_path / "ascocid.sqlite")
    index = IndexMemoire(tmp_path / "index")
    index.construire(["ascocid://chunk/existant"], [[1.0, 0.0, 0.0, 0.0]])
    yield Bibliotheque(tmp_path / "documents"), magasin, index
    magasin.cx.close()


# ── Lecture ──────────────────────────────────────────────────────────────────


def test_formats_acceptes_et_refuses():
    assert format_de("Guide.PDF") == "pdf"
    assert format_de("note.docx") == "docx"
    with pytest.raises(DocumentIllisible, match=r"\.doc"):
        format_de("ancien.doc")
    with pytest.raises(DocumentIllisible, match="PDF et Word"):
        format_de("tableur.xlsx")


def test_pdf_sections_titres_et_entetes_retires():
    entete = ("IFPC - Guide technique", 9)
    octets = pdf([
        [entete, ("Fermentation", 18), (PARAGRAPHE * 2, 11), ("Page 1", 9)],
        [entete, ("Clarification", 18),
         ("La clarification par défécation repose sur la formation d'un chapeau "
          "brun à la surface du moût, qui emprisonne les particules.", 11), ("Page 2", 9)],
        [entete, ("Le soutirage intervient lorsque la densité a suffisamment baissé "
                  "et que le chapeau est bien formé, sans attendre sa retombée.", 11),
         ("Page 3", 9)],
    ])
    lu = lire(octets, "guide.pdf")
    titres = [s.titre for s in lu.sections]
    assert titres[:2] == ["Fermentation", "Clarification"]
    assert lu.sections[1].page == 2
    texte = " ".join(p for s in lu.sections for p in s.paragraphes)
    # L'en-tête et le pied répétés sur chaque page ne sont pas du contenu.
    assert "Guide technique" not in texte
    assert "Page " not in texte
    assert "10 et 14 degrés" in texte
    assert "soutirage" in texte


def test_phrase_repetee_dans_le_corps_conservee():
    avertissement = "Toujours mesurer la densité avant le soutirage, sans exception."
    octets = pdf([
        [("Titre du guide", 9), (PARAGRAPHE, 11), (avertissement, 11),
         (f"Suite propre à la page {n}, distincte des autres pages du guide.", 11),
         (f"Pied {n}", 9)]
        for n in range(1, 4)
    ])
    texte = " ".join(p for s in lire(octets, "g.pdf").sections for p in s.paragraphes)
    assert texte.count("Toujours mesurer la densité") == 3


def test_pdf_sans_texte_refuse_comme_scan():
    import pymupdf

    document = pymupdf.open()
    document.new_page()
    with pytest.raises(DocumentIllisible, match="scanné"):
        lire(document.tobytes(), "scan.pdf")


def test_fichier_corrompu_refuse():
    with pytest.raises(DocumentIllisible, match="PDF lisible"):
        lire(b"ceci n'est pas un pdf", "faux.pdf")
    with pytest.raises(DocumentIllisible, match="Word"):
        lire(b"ceci n'est pas un docx", "faux.docx")


def test_docx_titres_paragraphes_et_tableaux_dans_l_ordre():
    lu = lire(docx([
        ("titre", "Conduite de la fermentation"),
        ("texte", PARAGRAPHE),
        ("tableau", "Étape;Température\nPrise de mousse;12 °C"),
        ("titre", "Conditionnement"),
        ("texte", PARAGRAPHE),
    ]), "note-technique.docx")
    assert [s.titre for s in lu.sections] == ["Conduite de la fermentation", "Conditionnement"]
    # Le tableau reste dans sa section : c'est souvent là qu'est la valeur cherchée.
    assert "Prise de mousse | 12 °C" in lu.sections[0].paragraphes
    assert lu.titre == "Conduite de la fermentation"


# ── Ajout, recherche, retrait ────────────────────────────────────────────────


def test_un_document_ajoute_est_trouvable_puis_introuvable_une_fois_retire(espace):
    biblio, magasin, index = espace
    doc, lu = biblio.enregistrer(
        "note.docx", docx([("titre", "Températures"), ("texte", PARAGRAPHE * 3)]),
        titre="Note fermentation", auteur="admin@ifpc.com")
    assert doc.idoc == PREMIER_IDOC and doc.etat == "en_cours"
    assert biblio.chemin(doc).exists()

    fini = biblio.integrer(doc, lu, magasin, EmbedderFactice(), index)
    assert fini is not None and fini.etat == "pret" and fini.passages > 0

    # Trouvable par les deux voies de la recherche hybride.
    trouves = magasin.recherche_lexicale('"aromes"', 10)
    assert trouves, "la recherche lexicale doit trouver le document"
    cles = list(magasin.chunks_par_cles([c for c, _ in trouves]))
    assert all(magasin.chunks_par_cles([c])[c]["fiche_idoc"] == doc.idoc for c in cles)
    assert index.taille == 1 + fini.passages
    fiche = magasin.cx.execute("SELECT * FROM fiche WHERE idoc=?", (doc.idoc,)).fetchone()
    assert fiche["type"] == TypeFiche.DOCUMENT.value and fiche["titre"] == "Note fermentation"

    retire = biblio.retirer(doc.idoc, magasin, index)
    assert retire is not None
    assert magasin.recherche_lexicale('"aromes"', 10) == []
    assert index.taille == 1                       # le corpus d'origine est intact
    assert not biblio.chemin(doc).exists()
    assert biblio.lister() == []


def test_doublon_refuse_sans_rien_ecrire(espace):
    biblio, _, _ = espace
    octets = docx([("texte", PARAGRAPHE * 2)])
    biblio.enregistrer("a.docx", octets)
    with pytest.raises(DocumentRefuse, match="déjà été ajouté"):
        biblio.enregistrer("copie.docx", octets)
    assert len(biblio.lister()) == 1


def test_refus_ne_laisse_aucune_trace(espace):
    biblio, _, _ = espace
    import pymupdf

    vide = pymupdf.open()
    vide.new_page()
    with pytest.raises(DocumentRefuse, match="scanné"):
        biblio.enregistrer("scan.pdf", vide.tobytes())
    with pytest.raises(DocumentRefuse, match="vide"):
        biblio.enregistrer("rien.pdf", b"")
    assert biblio.lister() == []
    assert not biblio.fichiers.exists() or not any(biblio.fichiers.iterdir())


def test_retrait_pendant_l_integration_n_est_pas_annule(espace):
    biblio, magasin, index = espace
    doc, lu = biblio.enregistrer("n.docx", docx([("texte", PARAGRAPHE * 2)]))
    biblio.retirer(doc.idoc, magasin, index)
    # L'intégration arrive après le retrait : elle ne doit rien ressusciter.
    assert biblio.integrer(doc, lu, magasin, EmbedderFactice(), index) is None
    assert magasin.cx.execute("SELECT count(*) FROM fiche").fetchone()[0] == 0


def test_reconstruction_du_corpus_conserve_les_documents(espace):
    biblio, magasin, index = espace
    doc, lu = biblio.enregistrer("n.docx", docx([("texte", PARAGRAPHE * 2)]))
    biblio.integrer(doc, lu, magasin, EmbedderFactice(), index)

    # Ce que fait `ingest extraire` : base vidée, AsCoCid réécrit…
    magasin.vider()
    magasin.ecrire(Corpus(fiches=[Fiche(idoc=1014, code="FICHE_X", type=TypeFiche.ARTICLE,
                                        titre="Fiche AsCoCid")]))
    assert reinjecter(biblio, magasin) == []
    titres = {r["titre"] for r in magasin.cx.execute("SELECT titre FROM fiche")}
    assert titres == {"Fiche AsCoCid", doc.titre}


# ── Index vectoriel à chaud ──────────────────────────────────────────────────


def test_index_ajout_remplacement_et_retrait(tmp_path):
    index = IndexMemoire(tmp_path)
    index.construire(["a", "b"], [[1, 0], [0, 1]])
    index.ajouter(["c", "a"], [[1, 1], [0, 1]])
    assert index.taille == 3
    assert index.chercher([0, 1], 1)[0][0] in {"a", "b"}
    index.retirer(["a", "inconnue"])
    assert index.taille == 2

    # Ce qui est écrit sur le disque est ce qui est en mémoire.
    relu = IndexMemoire(tmp_path)
    relu.charger()
    assert relu.taille == 2
    assert {c for c, _ in relu.chercher([1, 1], 5)} == {"b", "c"}


def test_index_absent_cree_au_premier_ajout(tmp_path):
    index = IndexMemoire(tmp_path / "neuf")
    index.retirer(["x"])                      # sans effet, sans erreur
    index.ajouter(["x"], [[1.0, 2.0]])
    assert index.taille == 1


# ── Présentation des sources ─────────────────────────────────────────────────


def test_source_document_renvoie_vers_son_fichier_et_non_vers_ascocid(espace):
    biblio, magasin, index = espace
    doc, lu = biblio.enregistrer("n.docx", docx([("texte", PARAGRAPHE * 2)]))
    biblio.integrer(doc, lu, magasin, EmbedderFactice(), index)
    magasin.ecrire(Corpus(fiches=[Fiche(idoc=1014, code="FICHE_X", type=TypeFiche.ARTICLE,
                                        titre="Fiche AsCoCid")]))

    provenances = magasin.provenances([doc.idoc, 1014])
    assert Pipeline._provenance(doc.idoc, provenances) == {
        "provenance": "document", "url": f"/api/ldc/document/{doc.idoc}/fichier"}
    assert Pipeline._provenance(1014, provenances) == {
        "provenance": "ascocid",
        "url": "https://ascocid.fr/ldc/view.php?id_document=1014"}
